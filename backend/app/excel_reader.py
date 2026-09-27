"""Excel 解析引擎。

设计原则：**不假设表结构**。
- 表头按别名规则自动映射到系统字段；映射不上的列原样保留为「动态列」，驱动前端表格自动渲染。
- 「金额」是可选字段：源表有价格类列则关联，没有就留空，绝不臆造默认值。
- 「商标编号」按字符串读取（18-19 位数字超过 JS 安全整数范围，必须防精度丢失）。
- 图片取自单元格浮动锚点，按锚点行号与数据行一一对应（本数据集为 TwoCellAnchor，列 B，1 图 1 行）。
"""
from __future__ import annotations

import csv
import hashlib
import io
import math
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# --------------------------------------------------------------------------- #
# 字段映射规则
# --------------------------------------------------------------------------- #
# 系统核心字段 ← Excel 表头别名
FIELD_ALIASES: dict[str, list[str]] = {
    "trademark_no": ["商标编号", "注册号", "注册编号", "商标号", "申请号", "申请编号", "商标注册号", "注册证号"],
    "name": ["商标名", "商标名称", "名称", "品牌名", "商标"],
    "category": ["类别", "分类", "注册类别", "商标类别", "商品类别"],
    "price": ["价格", "金额", "售价", "报价", "转让价", "标价", "转让价格", "价格元", "金额元"],
    "products": ["产品/服务", "产品服务", "核定使用商品", "核定使用商品/服务项目", "商品/服务", "服务项目", "产品", "核定商品"],
    "groups": ["群组", "类似群", "类似群组", "类似群号", "群组类似群"],
    "registration_date": ["注册日期", "注册时间", "申请日期", "注册日", "核准注册日期"],
    "expiry_date": ["有效期至", "有效期", "专用权期限至", "到期日期", "到期日", "专用期限至"],
    "application_count": ["申请量", "申请数量", "申请件数", "申请数", "申请次数"],
    "ai_description": ["ai释义", "释义", "ai解读", "商标释义", "含义", "ai分析", "品牌释义"],
    "legal_status": ["法律状态", "商标状态", "状态"],
    "remark": ["备注", "说明", "附注", "备注说明", "其他"],
    "image": ["图样", "图片", "商标图样", "商标图片", "商标logo", "logo", "图样图片", "商标图"],
}

# 明确忽略、不入库的列
IGNORE_HEADERS = {"序号", "no", "no.", "id", "序列号", "行号", "编号序号", "index"}

CORE_LABELS: dict[str, str] = {
    "serial_no": "唯一编号",
    "trademark_no": "商标编号",
    "name": "商标名",
    "category": "类别",
    "price": "金额",
    "products": "产品/服务",
    "groups": "群组",
    "registration_date": "注册日期",
    "expiry_date": "有效期至",
    "application_count": "申请量",
    "ai_description": "AI释义",
    "legal_status": "法律状态",
    "remark": "备注",
    "image": "图样",
}

# 字符串匹配优先级：越具体的字段先匹配，避免「商标编号」被「商标」误吞
_SUBSTRING_PRIORITY = [
    "trademark_no", "expiry_date", "registration_date", "application_count",
    "ai_description", "price", "image", "groups", "products", "legal_status",
    "category", "remark", "name",
]

_DATE_PATTERNS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y%m%d")
_PRICE_CLEAN = re.compile(r"[￥¥,，\s元人民币]")
_INT_CLEAN = re.compile(r"[^\d]")


def normalize_header(raw: Any, idx: int) -> tuple[str, str]:
    """返回 (display_label, key)。key 稳定、唯一，用于动态列。"""
    label = "" if raw is None else str(raw).replace("\n", " ").strip()
    label = re.sub(r"\s+", " ", label)
    if not label:
        label = f"列{idx + 1}"
    key = label.lower().replace(" ", "").replace("/", "_").replace("（", "(").replace("）", ")")
    key = re.sub(r"[()【】\[\]]", "", key)
    return label, key


def match_core_field(label: str, key: str) -> str | None:
    """把表头映射到系统核心字段；返回 None 表示这是动态列。"""
    flat_label = label.lower().replace(" ", "").replace("/", "").replace("（", "").replace("）", "").replace("(", "").replace(")", "")
    for f, aliases in FIELD_ALIASES.items():
        for a in aliases:
            if flat_label == a.lower().replace("/", "").replace("(", "").replace(")", ""):
                return f
    # 二次机会：包含匹配（按优先级从具体到宽泛）
    for f in _SUBSTRING_PRIORITY:
        for a in FIELD_ALIASES[f]:
            token = a.lower().replace("/", "").replace("(", "").replace(")", "")
            if len(token) >= 2 and token in flat_label:
                return f
    return None


# --------------------------------------------------------------------------- #
# 数据结构
# --------------------------------------------------------------------------- #
@dataclass
class ColumnInfo:
    index: int                       # 0 基列序号
    header: str                      # 原表头
    key: str                         # 动态列 key
    mapped_field: str | None         # 命中的核心字段，None = 动态列
    data_type: str = "text"          # text | number | date | price | image
    samples: list[str] = field(default_factory=list)
    non_empty: int = 0

    @property
    def label(self) -> str:
        return CORE_LABELS.get(self.mapped_field or "", self.header)


@dataclass
class ImageRef:
    row: int                         # Excel 行号（1 基）
    filename: str
    rel_url: str
    size: int
    fmt: str


@dataclass
class ParsedSheet:
    filename: str
    sheet_name: str
    columns: list[ColumnInfo]
    rows: list[dict[str, Any]]       # 每行：{"_excel_row": n, "_cells": {key: value}, ...}
    images: dict[int, list[ImageRef]]  # Excel 行号 -> 图片
    warnings: list[str]


# --------------------------------------------------------------------------- #
# 值清洗
# --------------------------------------------------------------------------- #
def _to_text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s or None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() or None


def clean_id_text(v: Any) -> str | None:
    """商标编号：必须按字符串处理，且去掉 Excel 数字格式留下的 .0"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s or None


def _to_date(v: Any) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    # 长文本一律不当作日期，避免「群组」这类 "0301；0302；…" 被误解析
    if not s or len(s) > 20:
        return None
    for fmt in _DATE_PATTERNS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    m = re.search(r"((?:19|20)\d{2})\D{0,2}(\d{1,2})\D{0,2}(\d{1,2})", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def _looks_like_date(v: Any) -> bool:
    """严格判断：只认真正的日期对象或标准日期字符串。"""
    if isinstance(v, (date, datetime)):
        return True
    if not isinstance(v, str):
        return False
    s = v.strip()
    if len(s) > 12:
        return False
    return any(_safe_strptime(s, fmt) for fmt in _DATE_PATTERNS)


def _safe_strptime(s: str, fmt: str) -> bool:
    try:
        datetime.strptime(s, fmt)
        return True
    except ValueError:
        return False


def _to_number(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return None
        return float(v)
    s = _PRICE_CLEAN.sub("", str(v))
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _to_int(v: Any) -> int | None:
    n = _to_number(v)
    return int(n) if n is not None else None


def _to_price(v: Any) -> float | None:
    n = _to_number(v)
    if n is None or n < 0:
        return None
    return round(n, 2)


def _to_dynamic(v: Any, dtype: str) -> Any:
    """动态列取值：按推断出的类型保留原始形态（数值仍是数值，日期仍是日期）。"""
    if v is None or v == "":
        return None
    if dtype == "number":
        if isinstance(v, str):
            s = v.strip()
            if s.startswith("0") and len(s) > 1:      # 保留前导零，如 0301
                return s
            n = _to_number(v)
            if n is None:
                return s
            if "." not in s and n.is_integer():
                return int(n)
            return n
        n = _to_number(v)
        return int(n) if n is not None and float(n).is_integer() else n
    if dtype == "date":
        return _to_date(v)
    return _to_text(v)


def _sample_str(v: Any) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)[:60]


def _infer_type(values: list[Any], mapped: str | None) -> str:
    """已映射到核心字段的列，类型由字段语义决定，绝不靠猜，避免「群组」被误判为日期。"""
    if mapped == "image":
        return "image"
    if mapped == "price":
        return "price"
    if mapped in ("registration_date", "expiry_date"):
        return "date"
    if mapped == "application_count":
        return "number"
    if mapped in ("trademark_no", "name", "products", "groups", "ai_description",
                  "remark", "legal_status"):
        return "text"
    # 未映射的动态列：从数据形态推断
    filled = [v for v in values if v not in (None, "")]
    if not filled:
        return "text"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in filled):
        return "number"
    if all(_looks_like_date(v) for v in filled[:10]):
        return "date"
    return "text"


# --------------------------------------------------------------------------- #
# 主解析流程
# --------------------------------------------------------------------------- #
def _extract_images(ws, dest_dir: Path, url_prefix: str) -> dict[int, list[ImageRef]]:
    """把工作表内嵌图片按锚点行号分组落盘，返回 {excel_row: [ImageRef]}。"""
    dest_dir.mkdir(parents=True, exist_ok=True)
    grouped: dict[int, list[ImageRef]] = {}
    raw_images = list(getattr(ws, "_images", []) or [])
    for i, im in enumerate(raw_images):
        try:
            anchor = im.anchor
            frm = getattr(anchor, "_from", None)
            if frm is None:
                continue
            excel_row = int(frm.row) + 1
            col = int(frm.col)
        except Exception:
            continue
        try:
            data = im._data()  # noqa: SLF001 - openpyxl 图片原始字节
        except Exception:
            continue
        if not data:
            continue
        fmt = (getattr(im, "format", None) or "png").lower()
        ext = "jpg" if fmt in ("jpeg", "jpg") else fmt
        digest = hashlib.md5(data).hexdigest()[:10]
        filename = f"r{excel_row}c{col}_{digest}.{ext}"
        target = dest_dir / filename
        if not target.exists():
            target.write_bytes(data)
        ref = ImageRef(
            row=excel_row,
            filename=filename,
            rel_url=f"{url_prefix}/{filename}",
            size=len(data),
            fmt=ext,
        )
        grouped.setdefault(excel_row, []).append(ref)
    for refs in grouped.values():
        refs.sort(key=lambda r: r.filename)
    return grouped


def parse_xlsx(path: Path, images_dir: Path, url_prefix: str, sheet_name: str | None = None) -> ParsedSheet:
    wb = load_workbook(path, data_only=True, read_only=False)
    warnings: list[str] = []
    try:
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.worksheets[0]
        if sheet_name and sheet_name not in wb.sheetnames:
            warnings.append(f"未找到工作表「{sheet_name}」，已改用「{ws.title}」")

        max_col = ws.max_column or 0
        max_row = ws.max_row or 0
        if max_row < 2:
            raise ValueError("文件只有表头或为空，没有可导入的数据行")

        # 1) 定位表头行：取前 5 行中非空单元格最多的一行
        header_row = 1
        best = -1
        for r in range(1, min(5, max_row) + 1):
            filled = sum(
                1 for c in range(1, max_col + 1)
                if ws.cell(row=r, column=c).value not in (None, "")
            )
            if filled > best:
                best, header_row = filled, r
        if header_row != 1:
            warnings.append(f"第 1 行不是有效表头，已自动使用第 {header_row} 行作为表头")

        # 2) 列识别
        columns: list[ColumnInfo] = []
        used_keys: set[str] = set()
        for c in range(1, max_col + 1):
            raw = ws.cell(row=header_row, column=c).value
            label, key = normalize_header(raw, c - 1)
            base_key, n = key, 2
            while key in used_keys:
                key, n = f"{base_key}_{n}", n + 1
            used_keys.add(key)

            values = [ws.cell(row=r, column=c).value for r in range(header_row + 1, max_row + 1)]
            non_empty = sum(1 for v in values if v not in (None, ""))
            if label.lower() in IGNORE_HEADERS and non_empty == 0:
                continue
            mapped = match_core_field(label, key)
            if label.lower() in IGNORE_HEADERS:
                mapped = None  # 显式忽略列：即使有值也不入库（如「序号」）
                if non_empty:
                    continue
            samples = [_sample_str(v) for v in values if v not in (None, "")][:3]
            dtype = _infer_type(values, mapped)
            columns.append(ColumnInfo(
                index=c - 1, header=label, key=key, mapped_field=mapped,
                data_type=dtype, samples=samples, non_empty=non_empty,
            ))

        # 3) 图片（锚点行号 1 基）
        images = _extract_images(ws, images_dir, url_prefix)
        image_anchor_ok = True
        if images:
            data_rows = set(range(header_row + 1, max_row + 1))
            outside = [r for r in images if r not in data_rows]
            if outside:
                image_anchor_ok = False
                warnings.append(
                    f"{len(outside)} 张图片的锚点行落在数据区之外，已按最接近的行兜底关联"
                )

        # 4) 逐行取值
        rows: list[dict[str, Any]] = []
        for r in range(header_row + 1, max_row + 1):
            cells: dict[str, Any] = {}
            empty = True
            for col in columns:
                v = ws.cell(row=r, column=col.index + 1).value
                f = col.mapped_field
                if f == "image":
                    continue  # 图样列不取单元格值，图片单独处理
                val: Any
                if f == "trademark_no":
                    val = clean_id_text(v)
                elif f == "category":
                    val = _to_int(v)
                elif f == "price":
                    val = _to_price(v)
                elif f == "application_count":
                    val = _to_int(v)
                elif f in ("registration_date", "expiry_date"):
                    val = _to_date(v)
                elif f in ("name", "products", "groups", "ai_description", "remark", "legal_status"):
                    val = _to_text(v)
                else:
                    val = _to_dynamic(v, col.data_type)   # 动态列
                if val is not None:
                    empty = False
                cells[col.key] = val
            if empty and r not in images:
                continue
            rows.append({"_excel_row": r, "_cells": cells})

        # 5) 图片兜底关联：锚点越界的图片挂到最近的数据行
        if images and not image_anchor_ok and rows:
            row_of = [x["_excel_row"] for x in rows]
            for r in list(images.keys()):
                if r in row_of:
                    continue
                nearest = min(row_of, key=lambda x: abs(x - r))
                images.setdefault(nearest, []).extend(images.pop(r))

        return ParsedSheet(
            filename=path.name,
            sheet_name=ws.title,
            columns=columns,
            rows=rows,
            images=images,
            warnings=warnings,
        )
    finally:
        wb.close()


def parse_csv(path: Path) -> ParsedSheet:
    """CSV 没有图片，按 xlsx 同样的规则识别列。"""
    text = None
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("无法识别 CSV 文件编码，请另存为 UTF-8 或 xlsx 后重试")

    reader = list(csv.reader(io.StringIO(text)))
    if len(reader) < 2:
        raise ValueError("文件只有表头或为空，没有可导入的数据行")
    header = reader[0]
    columns: list[ColumnInfo] = []
    used: set[str] = set()
    for i, raw in enumerate(header):
        label, key = normalize_header(raw, i)
        base, n = key, 2
        while key in used:
            key, n = f"{base}_{n}", n + 1
        used.add(key)
        if label.lower() in IGNORE_HEADERS:
            continue
        values = [row[i] if i < len(row) else None for row in reader[1:]]
        non_empty = sum(1 for v in values if v not in (None, ""))
        mapped = match_core_field(label, key)
        columns.append(ColumnInfo(
            index=i, header=label, key=key, mapped_field=mapped,
            data_type=_infer_type(values, mapped),
            samples=[_sample_str(v) for v in values if v not in (None, "")][:3],
            non_empty=non_empty,
        ))
    rows: list[dict[str, Any]] = []
    for ri, raw in enumerate(reader[1:], start=2):
        cells: dict[str, Any] = {}
        empty = True
        for col in columns:
            v = raw[col.index] if col.index < len(raw) else None
            f = col.mapped_field
            if f == "trademark_no":
                val = clean_id_text(v)
            elif f in ("category", "application_count"):
                val = _to_int(v)
            elif f == "price":
                val = _to_price(v)
            elif f in ("registration_date", "expiry_date"):
                val = _to_date(v)
            elif f in ("name", "products", "groups", "ai_description", "remark", "legal_status"):
                val = _to_text(v)
            else:
                val = _to_dynamic(v, col.data_type)
            if val is not None:
                empty = False
            cells[col.key] = val
        if empty:
            continue
        rows.append({"_excel_row": ri, "_cells": cells})
    return ParsedSheet(path.name, "(CSV)", columns, rows, {}, [])


def parsed_to_json(parsed: ParsedSheet) -> dict:
    """落盘快照：日期统一转 ISO 字符串（JSON 不支持 date 对象）。"""
    date_keys = {c.key for c in parsed.columns if c.data_type == "date"}
    rows = []
    for row in parsed.rows:
        cells = {}
        for k, v in row["_cells"].items():
            if isinstance(v, (date, datetime)):
                cells[k] = v.isoformat()[:10]
            elif isinstance(v, float) and (v != v or v in (float("inf"), float("-inf"))):
                cells[k] = None
            else:
                cells[k] = v
        rows.append({"_excel_row": row["_excel_row"], "_cells": cells})
    return {
        "filename": parsed.filename,
        "sheet_name": parsed.sheet_name,
        "warnings": parsed.warnings,
        "columns": [asdict(c) for c in parsed.columns],
        "rows": rows,
        "date_keys": sorted(date_keys),
        "images": {str(k): [asdict(i) for i in v] for k, v in parsed.images.items()},
    }


def parsed_from_json(payload: dict) -> ParsedSheet:
    date_keys = set(payload.get("date_keys") or [])
    rows = []
    for row in payload.get("rows", []):
        cells = dict(row.get("_cells") or {})
        for k in date_keys & cells.keys():
            if isinstance(cells[k], str):
                cells[k] = _to_date(cells[k])
        rows.append({"_excel_row": row["_excel_row"], "_cells": cells})
    return ParsedSheet(
        filename=payload.get("filename", ""),
        sheet_name=payload.get("sheet_name", ""),
        columns=[ColumnInfo(**c) for c in payload.get("columns", [])],
        rows=rows,
        images={
            int(k): [ImageRef(**i) for i in v]
            for k, v in (payload.get("images") or {}).items()
        },
        warnings=payload.get("warnings", []),
    )