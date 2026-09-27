"""批量导入 API。

流程：上传（可多文件）→ 解析分析 → 列映射确认 → 预览 → 落库（后台任务 + 进度轮询）→ 结果。
解析阶段不写任何商标数据，误传文件不会污染库。
"""
from __future__ import annotations

import io
import json
import shutil
import uuid
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import desc, select
from sqlalchemy.orm import Session

from ..config import ALLOWED_EXCEL_EXT, MAX_UPLOAD_MB, UPLOAD_DIR
from ..database import get_db, new_session
from ..deps import current_admin
from ..importer import (
    IMPORT_ROOT, analyze_file, auto_mapping, build_preview, cleanup_batch_files, run_import,
)
from ..models import Admin, ImportBatch
from ..schemas import ImportCommitIn

router = APIRouter(prefix="/api/admin/imports", tags=["admin-imports"])

RAW_DIR = IMPORT_ROOT / "raw"


def _batch_dict(b: ImportBatch) -> dict:
    cols = b.columns_json or []
    return {
        "id": b.id,
        "filename": b.filename,
        "sheet_name": b.sheet_name,
        "total_rows": b.total_rows,
        "image_count": b.image_count,
        "status": b.status,
        "progress": b.progress,
        "success_count": b.success_count,
        "updated_count": b.updated_count,
        "skipped_count": b.skipped_count,
        "failed_count": b.failed_count,
        "message": b.message,
        "created_at": b.created_at.strftime("%Y-%m-%d %H:%M") if b.created_at else None,
        "finished_at": b.finished_at.strftime("%Y-%m-%d %H:%M") if b.finished_at else None,
        "column_count": len(cols),
        "columns": cols,
        "errors": (b.errors_json or [])[:200],
        "options": b.options_json or {},
    }


@router.post("/upload")
async def upload_files(
    background_tasks: BackgroundTasks,
    files: list[UploadFile] = File(...),
    db: Session = Depends(get_db),
    admin: Admin = Depends(current_admin),
):
    """支持一次选多个文件，逐个解析成独立批次。"""
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    results, failed = [], []
    for f in files:
        name = f.filename or "unnamed"
        ext = Path(name).suffix.lower()
        if ext not in ALLOWED_EXCEL_EXT:
            failed.append({"filename": name, "reason": "仅支持 .xlsx / .xlsm / .csv"})
            continue
        raw = await f.read()
        if len(raw) > MAX_UPLOAD_MB * 1024 * 1024:
            failed.append({"filename": name, "reason": f"文件超过 {MAX_UPLOAD_MB}MB 限制"})
            continue
        stored = RAW_DIR / f"{uuid.uuid4().hex[:8]}_{name}"
        stored.write_bytes(raw)

        batch = ImportBatch(filename=name, stored_path=str(stored), status="analyzing",
                            created_by=admin.id)
        db.add(batch)
        db.commit()
        db.refresh(batch)
        try:
            analyze_file(db, stored, batch.id, admin.id)
            results.append(_batch_dict(batch))
        except Exception as exc:  # noqa: BLE001
            batch.status = "failed"
            batch.message = f"解析失败：{exc}"
            db.commit()
            failed.append({"filename": name, "reason": str(exc)})
    return {"batches": results, "failed": failed}


@router.get("/history")
def history(limit: int = 50, db: Session = Depends(get_db), admin: Admin = Depends(current_admin)):
    rows = db.execute(
        select(ImportBatch).order_by(desc(ImportBatch.created_at)).limit(limit)
    ).scalars().all()
    return {"items": [_batch_dict(b) for b in rows]}


@router.get("/{batch_id}/preview")
def preview(batch_id: int, limit: int = 20, db: Session = Depends(get_db),
            admin: Admin = Depends(current_admin)):
    try:
        data = build_preview(db, batch_id, limit=limit)
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc
    # 附加自动映射建议，前端可一键采用
    from ..excel_reader import parsed_from_json
    from ..importer import _snapshot_path

    snap = _snapshot_path(batch_id)
    parsed = parsed_from_json(json.loads(snap.read_text(encoding="utf-8")))
    data["auto_mapping"] = auto_mapping(parsed)
    return data


@router.get("/{batch_id}/status")
def status(batch_id: int, db: Session = Depends(get_db), admin: Admin = Depends(current_admin)):
    b = db.get(ImportBatch, batch_id)
    if not b:
        raise HTTPException(404, "批次不存在")
    return _batch_dict(b)


@router.post("/commit")
def commit(payload: ImportCommitIn, background_tasks: BackgroundTasks,
           db: Session = Depends(get_db), admin: Admin = Depends(current_admin)):
    b = db.get(ImportBatch, payload.batch_id)
    if not b:
        raise HTTPException(404, "批次不存在")
    if b.status == "running":
        raise HTTPException(400, "该批次正在导入中，请稍候")
    if b.status == "done":
        raise HTTPException(400, "该批次已导入完成，如需再次导入请重新上传文件")

    options = {
        "price_mode": payload.price_mode,
        "fixed_price": payload.fixed_price,
        "status": payload.status,
        "import_mode": payload.import_mode,
        "overwrite_images": payload.overwrite_images,
        "is_featured": payload.is_featured,
        "default_category": payload.default_category,
    }
    background_tasks.add_task(run_import, payload.batch_id, payload.mapping, options, new_session)
    return {"ok": True, "batch_id": payload.batch_id, "status": "running"}


@router.delete("/{batch_id}")
def delete_batch(batch_id: int, purge_data: bool = False, db: Session = Depends(get_db),
                 admin: Admin = Depends(current_admin)):
    b = db.get(ImportBatch, batch_id)
    if not b:
        raise HTTPException(404, "批次不存在")
    if b.status == "running":
        raise HTTPException(400, "导入进行中，无法删除批次")
    from ..models import Trademark

    deleted = 0
    if purge_data:
        from ..models import TrademarkImage
        ids = [r for r in db.execute(
            select(Trademark.id).where(Trademark.import_batch_id == batch_id)
        ).scalars()]
        if ids:
            db.query(TrademarkImage).filter(TrademarkImage.trademark_id.in_(ids)).delete(
                synchronize_session=False)
            deleted = db.query(Trademark).filter(Trademark.id.in_(ids)).delete(
                synchronize_session=False)
    db.delete(b)
    db.commit()
    cleanup_batch_files(batch_id, keep_images=not purge_data)
    return {"ok": True, "deleted_trademarks": deleted}


@router.get("/template")
def download_template():
    """标准导入模板：只放系统能识别的标准列名，并标注「金额」为可选。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from ..excel_reader import CORE_LABELS

    wb = Workbook()
    ws = wb.active
    ws.title = "商标导入模板"
    head = ["序号", "图样", "商标名", "类别", "价格", "产品/服务", "群组",
            "注册日期", "申请量", "商标编号", "AI释义", "备注"]
    note = ["系统自动生成，可留空", "把图片插入本列", "必填", "如 29/31/35", "可选，留空即未定价",
            "核定使用商品", "如 0301；0302", "格式 2020-08-07", "数字", "官方注册号，用于去重",
            "文本", "同名多类等信息"]
    fill = PatternFill("solid", fgColor="0F172A")
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = fill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center")
    for i, t in enumerate(note, start=1):
        c = ws.cell(row=2, column=i, value=t)
        c.font = Font(color="0369A1", size=9)
    sample = [1, "", "示例商标", 29, 1980, "加工过的坚果；肉；蛋", "2901；2902",
              "2020-08-07", 104, "711386408046645262", "示例释义文本", "同名多类:29;30;"]
    for i, v in enumerate(sample, start=1):
        ws.cell(row=3, column=i, value=v)
    ws.freeze_panes = "A2"
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["F"].width = 40
    ws.column_dimensions["J"].width = 22
    ws.column_dimensions["K"].width = 30

    info = wb.create_sheet("字段说明")
    info.append(["列名", "系统字段", "是否必填", "说明"])
    info["A1"].font = Font(bold=True)
    guide = [
        ("图样", "图片", "否", "图片直接插入单元格（浮动图片），系统按行自动关联，无需写文件名"),
        ("商标名", "name", "是", "缺失该列的行会被跳过并记入失败清单"),
        ("商标编号", "trademark_no", "建议", "官方注册号；重复时用于识别已存在的记录，不会覆盖「唯一编号」"),
        ("价格", "price", "否", "留空 = 该商标不显示价格，前台展示「面议」；也可导入后批量改价"),
        ("注册日期", "registration_date", "否", "系统按注册日期 +10 年自动推算「有效期至」"),
        ("日期格式", "-", "-", "支持 2020-08-07 / 2020/08/07 / 2020年8月7日 / 20200807"),
        ("新增列", "extra", "否", "模板之外的自定义列会被系统原样接收，并自动出现在后台表格中"),
    ]
    for row in guide:
        info.append(list(row))
    for col, w in zip("ABCD", (14, 16, 10, 70)):
        info.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from urllib.parse import quote
    fn = "商标批量导入模板.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fn)}"},
    )