"""报价单：生成分享链接（可加密码）、有效期、公开访问、导出 Excel。"""
from __future__ import annotations

import io
import secrets
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..database import get_db
from ..deps import current_user
from ..models import Quote, QuoteItem, Trademark, User
from ..schemas import QuoteCreateIn
from ..settings_store import get_all, get_public

router = APIRouter(prefix="/api/quotes", tags=["quotes"])


def _quote_no(db: Session) -> str:
    day = datetime.now().strftime("%Y%m%d")
    n = db.execute(select(Quote.id).order_by(Quote.id.desc()).limit(1)).scalar() or 0
    return f"Q{day}-{n + 1:04d}"


def _quote_dict(q: Quote, with_items: bool = True) -> dict:
    data = {
        "id": q.id,
        "quote_no": q.quote_no,
        "token": q.token,
        "title": q.title,
        "customer_name": q.customer_name,
        "contact_phone": q.contact_phone,
        "remark": q.remark,
        "total_original": float(q.total_original or 0),
        "total_quote": float(q.total_quote or 0),
        "has_password": bool(q.password),
        "expire_at": q.expire_at.strftime("%Y-%m-%d") if q.expire_at else None,
        "status": _effective_status(q),
        "view_count": q.view_count,
        "created_at": q.created_at.strftime("%Y-%m-%d %H:%M") if q.created_at else None,
        "item_count": len(q.items or []),
    }
    if with_items:
        data["items"] = [{
            "trademark_id": it.trademark_id,
            "serial_no": it.serial_no,
            "trademark_no": it.trademark_no,
            "name": it.trademark_name,
            "category": it.category,
            "original_price": float(it.original_price) if it.original_price is not None else None,
            "quote_price": float(it.quote_price) if it.quote_price is not None else None,
            "image": it.image_url,
        } for it in (q.items or [])]
    return data


def _effective_status(q: Quote) -> str:
    if q.status == "cancelled":
        return "cancelled"
    if q.expire_at and q.expire_at < datetime.now():
        return "expired"
    return "active"


@router.post("")
def create_quote(payload: QuoteCreateIn, user: User = Depends(current_user),
                 db: Session = Depends(get_db)):
    if not payload.items:
        raise HTTPException(400, "请先选择要报价的商标")
    cfg = get_all(db)
    days = payload.expire_days or int(cfg.get("quote_default_days", "7"))

    tm_ids = [i.trademark_id for i in payload.items]
    tms = {t.id: t for t in db.execute(
        select(Trademark).where(Trademark.id.in_(tm_ids)).options(selectinload(Trademark.images))
    ).scalars()}

    quote = Quote(
        quote_no=_quote_no(db),
        token=secrets.token_urlsafe(24).replace("-", "").replace("_", "")[:32],
        user_id=user.id,
        title=payload.title or "商标报价单",
        customer_name=payload.customer_name,
        contact_phone=payload.contact_phone or user.phone,
        remark=payload.remark,
        password=(payload.password or "").strip() or None,
        expire_at=datetime.now() + timedelta(days=max(1, min(days, 30))),
    )
    total_original = total_quote = 0.0
    for item in payload.items:
        tm = tms.get(item.trademark_id)
        if not tm:
            continue
        primary = next((im.url for im in (tm.images or []) if im.is_primary), None)
        if not primary and tm.images:
            primary = tm.images[0].url
        op = float(tm.price) if tm.price is not None else None
        qp = float(item.quote_price) if item.quote_price is not None else op
        total_original += op or 0
        total_quote += qp or 0
        quote.items.append(QuoteItem(
            trademark_id=tm.id, serial_no=tm.serial_no, trademark_no=tm.trademark_no,
            trademark_name=tm.name, category=tm.category, image_url=primary,
            original_price=op, quote_price=qp,
        ))
        tm.quote_count = (tm.quote_count or 0) + 1
    if not quote.items:
        raise HTTPException(400, "所选商标不存在或已下架")
    quote.total_original = round(total_original, 2)
    quote.total_quote = round(total_quote, 2)
    db.add(quote)
    db.commit()
    db.refresh(quote)
    return {"ok": True, "quote": _quote_dict(quote), "share_path": f"/quote/{quote.token}"}


@router.get("")
def my_quotes(user: User = Depends(current_user), db: Session = Depends(get_db)):
    rows = db.execute(
        select(Quote).where(Quote.user_id == user.id)
        .options(selectinload(Quote.items)).order_by(Quote.created_at.desc())
    ).scalars().all()
    return {"items": [_quote_dict(q, with_items=False) for q in rows]}


@router.delete("/{quote_id}")
def delete_quote(quote_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    q = db.get(Quote, quote_id)
    if not q or q.user_id != user.id:
        raise HTTPException(404, "报价单不存在")
    db.delete(q)
    db.commit()
    return {"ok": True}


@router.get("/share/{token}")
def share(token: str, password: str | None = None, db: Session = Depends(get_db)):
    q = db.execute(
        select(Quote).where(Quote.token == token).options(selectinload(Quote.items))
    ).scalar_one_or_none()
    if not q:
        raise HTTPException(404, "报价单不存在或链接已失效")
    status = _effective_status(q)
    if status == "cancelled":
        raise HTTPException(410, "该报价单已作废")
    if q.password and (password or "").strip() != q.password:
        return {"need_password": True, "title": q.title,
                "status": status, "expired": status == "expired"}

    q.view_count = (q.view_count or 0) + 1
    q.last_view_at = datetime.now()
    db.commit()
    return {
        "need_password": False,
        "expired": status == "expired",
        "quote": _quote_dict(q),
        "config": get_public(db),
    }


@router.get("/share/{token}/export")
def export_quote(token: str, password: str | None = None, db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from urllib.parse import quote as urlquote

    q = db.execute(
        select(Quote).where(Quote.token == token).options(selectinload(Quote.items))
    ).scalar_one_or_none()
    if not q:
        raise HTTPException(404, "报价单不存在")
    if q.password and (password or "").strip() != q.password:
        raise HTTPException(403, "访问密码不正确")

    cfg = get_all(db)
    wb = Workbook()
    ws = wb.active
    ws.title = "商标报价单"
    ws.merge_cells("A1:G1")
    title = ws["A1"]
    title.value = q.title
    title.font = Font(size=16, bold=True, color="0F172A")
    title.alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:G2")
    sub = ws["A2"]
    sub.value = (f"报价单号：{q.quote_no}    生成日期：{q.created_at.strftime('%Y-%m-%d')}    "
                 f"有效期至：{q.expire_at.strftime('%Y-%m-%d') if q.expire_at else '长期'}    "
                 f"客户：{q.customer_name or '-'}")
    sub.font = Font(size=10, color="475569")
    sub.alignment = Alignment(horizontal="center")

    head = ["序号", "商标名", "类别", "商标编号", "原价", "报价", "小计"]
    fill = PatternFill("solid", fgColor="0F172A")
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.fill = fill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center")
    r = 4
    for i, it in enumerate(q.items, start=1):
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=it.trademark_name)
        ws.cell(row=r, column=3, value=it.category)
        c = ws.cell(row=r, column=4, value=it.trademark_no)
        c.number_format = "@"
        ws.cell(row=r, column=5, value=float(it.original_price) if it.original_price is not None else None)
        ws.cell(row=r, column=6, value=float(it.quote_price) if it.quote_price is not None else None)
        ws.cell(row=r, column=7, value=f"=F{r}*1")
        r += 1
    total_row = r
    ws.cell(row=total_row, column=5, value="合计").font = Font(bold=True)
    ws.cell(row=total_row, column=6, value=float(q.total_quote)).font = Font(bold=True)
    ws.cell(row=total_row, column=7, value=f"=SUM(G4:G{r - 1})").font = Font(bold=True)

    note_row = total_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    n = ws.cell(row=note_row, column=1,
                value=f"备注：{q.remark or '无'}\n客服微信：{cfg.get('service_wechat') or '-'}    "
                      f"电话：{cfg.get('contact_phone') or '-'}\n{cfg.get('copyright') or ''}")
    n.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[note_row].height = 60

    for col, w in zip("ABCDEFG", (6, 18, 8, 22, 12, 12, 12)):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"{q.quote_no}_商标报价单.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urlquote(fn)}"},
    )